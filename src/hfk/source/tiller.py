"""Read Tiller workbooks into the normalized ledger model."""

from __future__ import annotations

import datetime as dt
import fnmatch
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

from hfk.ledger import BalanceRow, Category, Ledger, Transaction, money

REQUIRED_COLUMNS = {
    "Transactions": ("Date", "Description", "Category", "Amount", "Account"),
    "Categories": ("Category", "Group", "Type"),
    "Balance History": ("Date", "Account", "Balance"),
}
CATEGORY_TYPES = frozenset({"Expense", "Income", "Transfer"})


class TillerFormatError(ValueError):
    """Raised when a workbook cannot be read as the expected Tiller format."""


def find_workbook(data_dir: str | Path, pattern: str = "Tiller*.xls*") -> Path:
    """Return the newest non-lock workbook matching ``pattern`` in ``data_dir``."""

    directory = Path(data_dir)
    try:
        entries = tuple(directory.iterdir())
    except OSError:
        entries = ()

    pattern_lower = pattern.lower()
    candidates = [
        entry
        for entry in entries
        if entry.is_file()
        and not entry.name.startswith("~$")
        and fnmatch.fnmatchcase(entry.name.lower(), pattern_lower)
    ]
    if not candidates:
        raise TillerFormatError(
            f"Missing Tiller workbook: directory={directory}, pattern={pattern}"
        )

    return max(candidates, key=lambda path: (_version_number(path.name), path.stat().st_mtime))


def read_tiller(path: str | Path) -> Ledger:
    """Read a Tiller workbook file without modifying it."""

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return Ledger(
            transactions=_read_transactions(workbook),
            categories=_read_categories(workbook),
            balances=_read_balances(workbook),
        )
    finally:
        workbook.close()


def load_ledger(path_or_dir: str | Path) -> Ledger:
    """Read a workbook path, or locate and read a workbook inside a directory."""

    path = Path(path_or_dir)
    if path.is_dir():
        path = find_workbook(path)
    return read_tiller(path)


def _read_transactions(workbook: Any) -> tuple[Transaction, ...]:
    rows = []
    for row_number, row, headers in _data_rows(workbook, "Transactions"):
        rows.append(
            Transaction(
                date=_date_value(row, headers, "Transactions", row_number, "Date"),
                description=_required_text(
                    row, headers, "Transactions", row_number, "Description"
                ),
                category=_required_text(row, headers, "Transactions", row_number, "Category"),
                amount=_money_value(row, headers, "Transactions", row_number, "Amount"),
                account=_required_text(row, headers, "Transactions", row_number, "Account"),
                transaction_id=_optional_text(row, headers, "Transaction ID"),
                account_id=_optional_text(row, headers, "Account ID"),
                institution=_optional_text(row, headers, "Institution"),
                full_description=_optional_text(row, headers, "Full Description"),
            )
        )
    return tuple(rows)


def _read_categories(workbook: Any) -> tuple[Category, ...]:
    rows = []
    for row_number, row, headers in _data_rows(workbook, "Categories"):
        category_type = _required_text(row, headers, "Categories", row_number, "Type")
        if category_type not in CATEGORY_TYPES:
            expected = ", ".join(sorted(CATEGORY_TYPES))
            raise TillerFormatError(
                "Invalid Tiller category type: "
                f"sheet=Categories, row={row_number}, column=Type, "
                f"value={category_type!r}, expected={expected}"
            )

        rows.append(
            Category(
                name=_required_text(row, headers, "Categories", row_number, "Category"),
                group=_required_text(row, headers, "Categories", row_number, "Group"),
                type=category_type,
                hidden=_truthy(_cell(row, headers, "Hide From Reports")),
            )
        )
    return tuple(rows)


def _read_balances(workbook: Any) -> tuple[BalanceRow, ...]:
    rows = []
    for row_number, row, headers in _data_rows(workbook, "Balance History"):
        rows.append(
            BalanceRow(
                date=_date_value(row, headers, "Balance History", row_number, "Date"),
                account=_required_text(row, headers, "Balance History", row_number, "Account"),
                balance=_optional_money_value(
                    row, headers, "Balance History", row_number, "Balance"
                ),
                account_id=_optional_text(row, headers, "Account ID"),
                institution=_optional_text(row, headers, "Institution"),
                account_type=_optional_text(row, headers, "Type"),
                account_class=_optional_text(row, headers, "Class"),
                account_status=_optional_text(row, headers, "Account Status"),
            )
        )
    return tuple(rows)


def _data_rows(workbook: Any, sheet_name: str) -> tuple[tuple[int, tuple[Any, ...], dict[str, int]], ...]:
    if sheet_name not in workbook.sheetnames:
        raise TillerFormatError(
            f"Missing required Tiller sheet: sheet={sheet_name}, "
            f"fix=restore the {sheet_name} sheet from your Tiller workbook"
        )

    sheet = workbook[sheet_name]
    row_iter = sheet.iter_rows(values_only=True)
    header_row = next(row_iter, ())
    headers = _headers(header_row)
    _require_columns(sheet_name, headers)
    return tuple(
        (row_number, tuple(row), headers)
        for row_number, row in enumerate(row_iter, start=2)
        if not _blank_row(row)
    )


def _headers(row: tuple[Any, ...]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for index, value in enumerate(row):
        if value is None or value == "":
            continue
        header = str(value)
        if header not in headers:
            headers[header] = index
    return headers


def _require_columns(sheet_name: str, headers: dict[str, int]) -> None:
    for column in REQUIRED_COLUMNS[sheet_name]:
        if column not in headers:
            raise TillerFormatError(
                f"Missing required Tiller column: sheet={sheet_name}, column={column}, "
                f"fix=restore the {column} column in the {sheet_name} sheet"
            )


def _cell(row: tuple[Any, ...], headers: dict[str, int], column: str) -> Any:
    index = headers.get(column)
    if index is None or index >= len(row):
        return None
    return row[index]


def _date_value(
    row: tuple[Any, ...],
    headers: dict[str, int],
    sheet_name: str,
    row_number: int,
    column: str,
) -> dt.date:
    value = _cell(row, headers, column)
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str) and value.strip():
        parsed = _parse_date_string(value.strip())
        if parsed is not None:
            return parsed

    raise TillerFormatError(
        "Unparseable Tiller date: "
        f"sheet={sheet_name}, row={row_number}, column={column}, value={value!r}"
    )


def _parse_date_string(value: str) -> dt.date | None:
    try:
        return dt.date.fromisoformat(value)
    except ValueError:
        pass

    try:
        return dt.datetime.fromisoformat(value).date()
    except ValueError:
        pass

    for pattern in (r"^(\d{1,2})/(\d{1,2})/(\d{4})$", r"^(\d{4})/(\d{1,2})/(\d{1,2})$"):
        match = re.match(pattern, value)
        if match is None:
            continue
        first, second, third = (int(group) for group in match.groups())
        if pattern.startswith("^(\\d{4})"):
            year, month, day = first, second, third
        else:
            month, day, year = first, second, third
        try:
            return dt.date(year, month, day)
        except ValueError:
            continue
    return None


def _money_value(
    row: tuple[Any, ...],
    headers: dict[str, int],
    sheet_name: str,
    row_number: int,
    column: str,
) -> Decimal:
    value = _cell(row, headers, column)
    try:
        return money(value)
    except ValueError as exc:
        raise TillerFormatError(
            "Unparseable Tiller amount: "
            f"sheet={sheet_name}, row={row_number}, column={column}, value={value!r}"
        ) from exc


def _optional_money_value(
    row: tuple[Any, ...],
    headers: dict[str, int],
    sheet_name: str,
    row_number: int,
    column: str,
) -> Decimal | None:
    value = _cell(row, headers, column)
    if value is None or value == "":
        return None
    try:
        return money(value)
    except ValueError as exc:
        raise TillerFormatError(
            "Unparseable Tiller amount: "
            f"sheet={sheet_name}, row={row_number}, column={column}, value={value!r}"
        ) from exc


def _required_text(
    row: tuple[Any, ...],
    headers: dict[str, int],
    sheet_name: str,
    row_number: int,
    column: str,
) -> str:
    value = _cell(row, headers, column)
    if value is None or value == "":
        raise TillerFormatError(
            f"Missing required Tiller value: sheet={sheet_name}, row={row_number}, "
            f"column={column}, fix=restore the cell value"
        )
    return str(value)


def _optional_text(row: tuple[Any, ...], headers: dict[str, int], column: str) -> str | None:
    value = _cell(row, headers, column)
    if value is None or value == "":
        return None
    return str(value)


def _blank_row(row: tuple[Any, ...]) -> bool:
    return all(value is None or value == "" for value in row)


def _truthy(value: Any) -> bool:
    if value is None or value == "":
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        return value != 0
    if isinstance(value, int):
        return value != 0

    text = str(value).strip().lower()
    if text in {"", "0", "false", "n", "no"}:
        return False
    if text in {"1", "true", "y", "yes", "x"}:
        return True
    return bool(text)


def _version_number(name: str) -> int:
    versions = [int(match) for match in re.findall(r"v(\d+)", name, flags=re.IGNORECASE)]
    if not versions:
        return -1
    return max(versions)
