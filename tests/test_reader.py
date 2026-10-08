from __future__ import annotations

import os
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from hfk.ledger import Ledger
from hfk.source import TillerFormatError, find_workbook, load_ledger, read_tiller
from hfk.source.fixture import build_fixture

REQUIRED_COLUMNS = [
    ("Transactions", "Date"),
    ("Transactions", "Description"),
    ("Transactions", "Category"),
    ("Transactions", "Amount"),
    ("Transactions", "Account"),
    ("Categories", "Category"),
    ("Categories", "Group"),
    ("Categories", "Type"),
    ("Balance History", "Date"),
    ("Balance History", "Account"),
    ("Balance History", "Balance"),
]


def test_reader_opens_workbook_read_only(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = build_fixture(tmp_path / "Tiller Test v1.xlsx")
    before = path.read_bytes()
    calls: list[dict[str, object]] = []
    real_load_workbook = openpyxl.load_workbook

    def recording_load_workbook(filename: object, **kwargs: object) -> Workbook:
        calls.append({"filename": filename, **kwargs})
        return real_load_workbook(filename, **kwargs)

    def fail_save(self: Workbook, filename: object | None = None) -> None:
        raise AssertionError(f"Workbook.save was called for {filename or self.path}")

    monkeypatch.setattr(openpyxl, "load_workbook", recording_load_workbook)
    monkeypatch.setattr(Workbook, "save", fail_save)

    ledger = read_tiller(path)

    assert isinstance(ledger, Ledger)
    assert calls == [{"filename": path, "read_only": True, "data_only": True}]
    assert path.read_bytes() == before


def test_reordered_headers_give_identical_ledgers(tmp_path: Path) -> None:
    plain = build_fixture(tmp_path / "plain.xlsx")
    shuffled = build_fixture(tmp_path / "shuffled.xlsx", shuffle_headers=True)

    assert read_tiller(shuffled) == read_tiller(plain)


def test_extra_columns_are_tolerated(tmp_path: Path) -> None:
    path = build_fixture(
        tmp_path / "fixture.xlsx",
        extra_columns=[
            ("Transactions", "Review Note"),
            ("Categories", "Owner Note"),
            ("Balance History", "Balance Note"),
        ],
    )

    ledger = read_tiller(path)

    assert len(ledger.transactions) > 0
    assert len(ledger.categories) > 0
    assert len(ledger.balances) > 0


def test_extra_sheets_are_tolerated(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    workbook = load_workbook(path)
    workbook.create_sheet("Notes")
    workbook.save(path)

    assert read_tiller(path).account_names()


@pytest.mark.parametrize(("sheet", "column"), REQUIRED_COLUMNS)
def test_missing_required_columns_raise_documented_message(
    tmp_path: Path, sheet: str, column: str
) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx", drop_columns=[(sheet, column)])

    with pytest.raises(TillerFormatError) as error:
        read_tiller(path)

    assert str(error.value) == (
        f"Missing required Tiller column: sheet={sheet}, column={column}, "
        f"fix=restore the {column} column in the {sheet} sheet"
    )


def test_missing_required_sheet_raises_documented_message(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    workbook = load_workbook(path)
    del workbook["Categories"]
    workbook.save(path)

    with pytest.raises(TillerFormatError) as error:
        read_tiller(path)

    assert str(error.value) == (
        "Missing required Tiller sheet: sheet=Categories, "
        "fix=restore the Categories sheet from your Tiller workbook"
    )


def test_invalid_category_type_names_sheet_row_column_and_value(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    _set_cell(path, "Categories", "Type", 2, "Spending")

    with pytest.raises(TillerFormatError) as error:
        read_tiller(path)

    message = str(error.value)
    assert "sheet=Categories" in message
    assert "row=2" in message
    assert "column=Type" in message
    assert "value='Spending'" in message


def test_invalid_transaction_date_names_sheet_row_column_and_value(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    _set_cell(path, "Transactions", "Date", 2, "not a date")

    with pytest.raises(TillerFormatError) as error:
        read_tiller(path)

    message = str(error.value)
    assert "sheet=Transactions" in message
    assert "row=2" in message
    assert "column=Date" in message
    assert "value='not a date'" in message


def test_invalid_transaction_amount_names_sheet_row_column_and_value(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    _set_cell(path, "Transactions", "Amount", 2, "not money")

    with pytest.raises(TillerFormatError) as error:
        read_tiller(path)

    message = str(error.value)
    assert "sheet=Transactions" in message
    assert "row=2" in message
    assert "column=Amount" in message
    assert "value='not money'" in message


def test_invalid_balance_amount_names_sheet_row_column_and_value(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    _set_cell(path, "Balance History", "Balance", 2, "not money")

    with pytest.raises(TillerFormatError) as error:
        read_tiller(path)

    message = str(error.value)
    assert "sheet=Balance History" in message
    assert "row=2" in message
    assert "column=Balance" in message
    assert "value='not money'" in message


def test_expenses_negative_income_positive_preserved(tmp_path: Path) -> None:
    ledger = read_tiller(build_fixture(tmp_path / "fixture.xlsx"))

    expense_amounts = [
        transaction.amount
        for transaction in ledger.transactions
        if ledger.category_type(transaction.category) == "Expense"
    ]
    income_amounts = [
        transaction.amount
        for transaction in ledger.transactions
        if ledger.category_type(transaction.category) == "Income"
    ]

    assert expense_amounts
    assert income_amounts
    assert all(amount < 0 for amount in expense_amounts)
    assert all(amount > 0 for amount in income_amounts)


def test_balance_gap_is_none_and_latest_balances_skip_it(tmp_path: Path) -> None:
    ledger = read_tiller(build_fixture(tmp_path / "fixture.xlsx", balance_gap=True))

    gaps = ledger.balance_gaps()
    latest = ledger.latest_balances()

    assert len(gaps) == 1
    assert gaps[0].balance is None
    assert latest[gaps[0].account].balance is not None


def test_blank_rows_are_skipped(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    ledger = read_tiller(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Transactions"]
    nonblank_data_rows = sum(
        1 for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)
    )

    assert len(ledger.transactions) == nonblank_data_rows


def test_optional_fields_are_mapped_and_hidden_truthy_is_true(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    _set_cell(path, "Categories", "Hide From Reports", 2, "TRUE")

    ledger = read_tiller(path)
    transaction = ledger.transactions[0]
    category = ledger.categories[0]
    balance = ledger.balances[0]

    assert transaction.transaction_id == "fx-0001"
    assert transaction.account_id == "acct-alex-checking"
    assert transaction.institution == "Example Bank"
    assert transaction.full_description == "Alex Paycheck example transaction"
    assert category.hidden is True
    assert balance.account_id == "acct-alex-checking"
    assert balance.institution == "Example Bank"
    assert balance.account_type == "Checking"
    assert balance.account_class == "Asset"
    assert balance.account_status == "Open"


def test_lock_file_is_ignored(tmp_path: Path) -> None:
    valid = build_fixture(tmp_path / "Tiller Test v1.xlsx")
    (tmp_path / "~$Tiller Test v2.xlsx").write_text("not a workbook", encoding="utf-8")

    assert find_workbook(tmp_path) == valid


def test_newest_version_chosen_numerically_before_mtime(tmp_path: Path) -> None:
    v2 = build_fixture(tmp_path / "Tiller Test v2.xlsx")
    v10 = build_fixture(tmp_path / "Tiller Test v10.xlsx")
    os.utime(v2, (200, 200))
    os.utime(v10, (100, 100))

    assert find_workbook(tmp_path) == v10


def test_same_version_uses_newest_mtime(tmp_path: Path) -> None:
    older = build_fixture(tmp_path / "Tiller Test v2.xlsx")
    newer = build_fixture(tmp_path / "Tiller Copy v2.xlsx")
    os.utime(older, (100, 100))
    os.utime(newer, (200, 200))

    assert find_workbook(tmp_path) == newer


def test_find_workbook_is_case_insensitive(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "tiller test V3.XLSX")

    assert find_workbook(tmp_path, pattern="TILLER*.xls*") == path


def test_no_match_error_names_directory_and_pattern(tmp_path: Path) -> None:
    with pytest.raises(TillerFormatError) as error:
        find_workbook(tmp_path, pattern="Tiller*.xls*")

    message = str(error.value)
    assert f"directory={tmp_path}" in message
    assert "pattern=Tiller*.xls*" in message


def test_load_ledger_on_directory_finds_workbook(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "Tiller Test v3.xlsx")

    assert load_ledger(tmp_path) == read_tiller(path)


def test_load_ledger_on_file_reads_that_file(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "example.xlsx")

    assert load_ledger(path) == read_tiller(path)


def test_default_fixture_transfer_pairs_net_to_zero(tmp_path: Path) -> None:
    ledger = read_tiller(build_fixture(tmp_path / "fixture.xlsx"))
    transfers = tuple(
        transaction
        for transaction in ledger.transactions
        if ledger.category_type(transaction.category) == "Transfer"
    )

    assert transfers
    assert Ledger.total(transfers) == Decimal("0.00")


def _set_cell(path: Path, sheet_name: str, header: str, row: int, value: object) -> None:
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    column = headers.index(header) + 1
    sheet.cell(row=row, column=column).value = value
    workbook.save(path)
