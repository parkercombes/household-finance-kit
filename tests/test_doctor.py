from __future__ import annotations

import os
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from hfk.engine import run_doctor
from hfk.source.fixture import build_fixture

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"
INSTANCE_EXAMPLE = EXAMPLES / "instance"
PYTHONPATH = str(ROOT / "src")


def make_instance(tmp_path: Path, **fixture_kwargs: Any) -> Path:
    for filename in ("household.yaml", "declarations.yaml", "plan.yaml"):
        shutil.copyfile(INSTANCE_EXAMPLE / filename, tmp_path / filename)
    build_fixture(tmp_path / "data" / "Tiller-Example.xlsx", **fixture_kwargs)
    return tmp_path


def run_hfk(*args: str) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONPATH"] = PYTHONPATH
    return subprocess.run(
        [sys.executable, "-m", "hfk", *args],
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )


def load_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    return data


def dump_yaml(path: Path, data: dict[str, Any]) -> None:
    path.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")


def fixture_workbook(home: Path) -> Path:
    return next((home / "data").glob("Tiller*.xlsx"))


def rename_first_transaction_category(home: Path, old: str, new: str) -> None:
    path = fixture_workbook(home)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Transactions"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        category_column = headers["Category"]
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=category_column).value == old:
                sheet.cell(row=row, column=category_column).value = new
                workbook.save(path)
                return
    finally:
        workbook.close()
    raise AssertionError(f"Category {old!r} not found in fixture workbook")


def orphan_diagnostics(report: Any) -> list[Any]:
    return [
        diagnostic
        for diagnostic in report.errors()
        if diagnostic.code.startswith("category-") or diagnostic.code.startswith("group-")
    ]


def test_example_instance_passes_doctor(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    report = run_doctor(home, today=date(2026, 10, 1))
    result = run_hfk("doctor", "--home", str(home), "--today", "2026-10-01")

    assert report.errors() == ()
    assert report.ok
    assert result.returncode == 0
    assert "doctor: 0 errors" in result.stdout
    assert "doctor: 0 errors, 0 warnings" in result.stdout
    assert result.stderr == ""


def test_example_instance_files_match_examples() -> None:
    pairs = {
        "household.yaml": "household.example.yaml",
        "declarations.yaml": "declarations.example.yaml",
        "plan.yaml": "plan.example.yaml",
    }
    for instance_name, example_name in pairs.items():
        assert (INSTANCE_EXAMPLE / instance_name).read_text(encoding="utf-8") == (
            EXAMPLES / example_name
        ).read_text(encoding="utf-8")


def test_missing_workbook_is_an_error(tmp_path: Path) -> None:
    for filename in ("household.yaml", "declarations.yaml", "plan.yaml"):
        shutil.copyfile(INSTANCE_EXAMPLE / filename, tmp_path / filename)

    report = run_doctor(tmp_path, today=date(2026, 10, 1))
    result = run_hfk("doctor", "--home", str(tmp_path), "--today", "2026-10-01")

    missing = [diagnostic for diagnostic in report.errors() if diagnostic.code == "workbook-missing"]
    assert len(missing) == 1
    assert "Tiller*.xls*" in missing[0].message
    assert result.returncode == 1
    assert "workbook-missing" not in result.stdout
    assert "Tiller*.xls*" in result.stdout


def test_config_errors_name_file_and_key(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    path = home / "household.yaml"
    data = load_yaml(path)
    del data["people"]
    dump_yaml(path, data)

    report = run_doctor(home, today=date(2026, 10, 1))

    config_errors = [diagnostic for diagnostic in report.errors() if diagnostic.code == "config-error"]
    assert any(
        diagnostic.where == "household.yaml" and "key=people" in diagnostic.message
        for diagnostic in config_errors
    )
    assert not any(diagnostic.where == "declarations.yaml" for diagnostic in config_errors)


def test_undeclared_account_is_an_error(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    del household["accounts"]["Savings"]
    dump_yaml(household_path, household)

    report = run_doctor(home, today=date(2026, 10, 1))
    undeclared = [
        diagnostic for diagnostic in report.errors() if diagnostic.code == "account-undeclared"
    ]
    assert any("Savings" in diagnostic.message for diagnostic in undeclared)

    declarations_path = home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["closed_accounts"] = {
        "Savings": {"closed": "2026-01-01", "note": "Closed in the example instance."}
    }
    dump_yaml(declarations_path, declarations)

    report = run_doctor(home, today=date(2026, 10, 1))
    assert not any(diagnostic.code == "account-undeclared" for diagnostic in report.errors())


def test_orphan_category_fails_doctor(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    rename_first_transaction_category(home, "Groceries", "Groceries (old)")

    report = run_doctor(home, today=date(2026, 10, 1))
    result = run_hfk("doctor", "--home", str(home), "--today", "2026-10-01")

    orphans = [diagnostic for diagnostic in report.errors() if diagnostic.code == "category-orphan"]
    assert len(orphans) == 1
    assert "Groceries (old)" in orphans[0].message
    assert "Categories sheet" in orphans[0].fix
    assert result.returncode == 1
    assert 'category "Groceries (old)" appears in transactions' in result.stdout


def test_unmapped_group_fails_doctor(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    declarations_path = home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    del declarations["groups"]["Subscriptions"]
    dump_yaml(declarations_path, declarations)

    report = run_doctor(home, today=date(2026, 10, 1))

    unmapped = [diagnostic for diagnostic in report.errors() if diagnostic.code == "group-unmapped"]
    assert len(unmapped) == 1
    assert "Subscriptions" in unmapped[0].message
    assert "Annual Membership" in unmapped[0].message

    declarations["non_spend_groups"].append("Subscriptions")
    dump_yaml(declarations_path, declarations)
    report = run_doctor(home, today=date(2026, 10, 1))
    assert not orphan_diagnostics(report)

    declarations["non_spend_groups"].remove("Subscriptions")
    declarations["overrides"]["Annual Membership"] = {
        "bucket": "Fixed Discretionary",
        "reason": "Annual memberships are optional in the example budget.",
    }
    dump_yaml(declarations_path, declarations)
    report = run_doctor(home, today=date(2026, 10, 1))
    assert not orphan_diagnostics(report)


def test_explicit_mapping_source_needs_every_category(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    declarations_path = home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["mapping_source"] = "explicit"
    del declarations["groups"]
    declarations["categories"] = {
        "Rent": "Fixed Essential",
        "Utilities": "Variable Essential",
        "Groceries": "Variable Essential",
        "Dining": "Variable Discretionary",
        "Fuel": "Variable Essential",
        "Annual Membership": "Fixed Discretionary",
    }
    dump_yaml(declarations_path, declarations)

    report = run_doctor(home, today=date(2026, 10, 1))
    orphans = orphan_diagnostics(report)

    assert len(orphans) == 1
    assert orphans[0].code == "category-unmapped"
    assert "Transit" in orphans[0].message


def test_income_and_transfer_categories_need_no_mapping(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    report = run_doctor(home, today=date(2026, 10, 1))

    assert not orphan_diagnostics(report)


def test_closed_account_liveness(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    declarations_path = home / "declarations.yaml"

    household = load_yaml(household_path)
    savings_account = household["accounts"].pop("Savings")
    dump_yaml(household_path, household)
    declarations = load_yaml(declarations_path)
    declarations["closed_accounts"] = {
        "Savings": {"closed": "2026-08-01", "note": "Closed in the example instance."}
    }
    dump_yaml(declarations_path, declarations)

    report = run_doctor(home, today=date(2026, 10, 1))
    result = run_hfk("doctor", "--home", str(home), "--today", "2026-10-01")

    active = [
        diagnostic for diagnostic in report.errors() if diagnostic.code == "closed-account-active"
    ]
    assert len(active) == 1
    assert "Savings" in active[0].message
    assert "2026-08-01" in active[0].message
    assert result.returncode == 1
    assert "Savings" in result.stdout
    assert "2026-08-01" in result.stdout

    declarations["closed_accounts"]["Savings"]["closed"] = "2026-12-31"
    declarations["closed_accounts"]["Old Card"] = {
        "closed": "2020-01-01",
        "note": "Closed before the example instance.",
    }
    dump_yaml(declarations_path, declarations)
    report = run_doctor(home, today=date(2026, 10, 1))

    assert report.ok
    assert not any(
        diagnostic.code in {"closed-account-active", "account-unused"}
        and "Savings" in diagnostic.message
        for diagnostic in report.diagnostics
    )
    assert not any("Old Card" in diagnostic.message for diagnostic in report.diagnostics)

    household["accounts"]["Savings"] = savings_account
    dump_yaml(household_path, household)
    report = run_doctor(home, today=date(2026, 10, 1))
    conflicts = [
        diagnostic
        for diagnostic in report.errors()
        if diagnostic.code == "account-closed-and-open"
    ]
    assert any("Savings" in diagnostic.message for diagnostic in conflicts)


def test_unused_declared_account_is_a_warning(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    path = home / "household.yaml"
    data = load_yaml(path)
    data["accounts"]["Old Card"] = {"owner": "alex"}
    dump_yaml(path, data)

    report = run_doctor(home, today=date(2026, 10, 1))

    unused = [diagnostic for diagnostic in report.warnings() if diagnostic.code == "account-unused"]
    assert report.ok
    assert any("Old Card" in diagnostic.message for diagnostic in unused)


def test_stale_data_is_a_warning(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    stale = run_doctor(home, today=date(2027, 1, 1))
    fresh = run_doctor(home, today=date(2026, 10, 1))

    assert any(diagnostic.code == "data-stale" for diagnostic in stale.warnings())
    assert not any(diagnostic.code == "data-stale" for diagnostic in fresh.warnings())


def test_history_start_checks(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    path = home / "household.yaml"

    data = load_yaml(path)
    del data["history_start"]
    dump_yaml(path, data)
    report = run_doctor(home, today=date(2026, 10, 1))
    assert any(diagnostic.code == "history-start-missing" for diagnostic in report.warnings())

    data = load_yaml(path)
    data["history_start"] = "2030-01-01"
    dump_yaml(path, data)
    report = run_doctor(home, today=date(2026, 10, 1))
    assert any(diagnostic.code == "history-start-after-data" for diagnostic in report.errors())

    data = load_yaml(path)
    data["history_start"] = "2026-07-01"
    dump_yaml(path, data)
    report = run_doctor(home, today=date(2026, 10, 1))
    clamps = [
        diagnostic for diagnostic in report.warnings() if diagnostic.code == "history-start-clamps"
    ]
    assert clamps
    count = int(clamps[0].message.split()[0])
    assert count > 0
