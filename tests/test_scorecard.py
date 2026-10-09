from __future__ import annotations

import os
import shutil
import subprocess
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from hfk.config import Month, load_declarations, load_household, load_plan
from hfk.engine import (
    BucketError,
    ScorecardError,
    allocate,
    build_scorecard,
    resolve_ledger,
    share_transaction,
)
from hfk.source import find_workbook, read_tiller
from hfk.source.fixture import build_fixture

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"
INSTANCE_EXAMPLE = EXAMPLES / "instance"
PYTHONPATH = str(ROOT / "src")


def make_instance(tmp_path: Path, **fixture_kwargs: Any) -> Path:
    tmp_path.mkdir(parents=True, exist_ok=True)
    for filename in ("household.yaml", "declarations.yaml", "plan.yaml"):
        shutil.copyfile(INSTANCE_EXAMPLE / filename, tmp_path / filename)
    build_fixture(tmp_path / "data" / "Tiller-Example.xlsx", **fixture_kwargs)
    return tmp_path


def run_hfk(*args: str) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONPATH"] = PYTHONPATH
    return subprocess.run(
        [sys.executable, "-m", "hfk", *args],
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )


def load_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    return data


def dump_yaml(path: Path, data: dict[str, Any]) -> None:
    path.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")


def fixture_workbook(home: Path) -> Path:
    return next((home / "data").glob("Tiller*.xlsx"))


def append_transaction(
    home: Path,
    *,
    day: int,
    description: str,
    category: str,
    amount: Decimal,
    account: str,
) -> None:
    transaction_date = date(2026, 9, day)
    path = fixture_workbook(home)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Transactions"]
        headers = [cell.value for cell in sheet[1]]
        account_id = account.lower().replace(" ", "-")
        values = {
            "Date": transaction_date,
            "Description": description,
            "Category": category,
            "Amount": amount,
            "Account": account,
            "Account #": None,
            "Institution": "Example Bank",
            "Month": date(2026, 9, 1),
            "Week": transaction_date,
            "Transaction ID": f"test-{day}-{description.lower().replace(' ', '-')}",
            "Account ID": f"acct-{account_id}",
            "Check Number": None,
            "Full Description": f"{description} example transaction",
            "Date Added": transaction_date,
            "Category Hint": category,
            "Categorized Date": transaction_date,
        }
        sheet.append([values.get(header) for header in headers])
        workbook.save(path)
    finally:
        workbook.close()


def rename_first_transaction_category(
    home: Path, old: str, new: str, *, month: str = "2026-09"
) -> None:
    path = fixture_workbook(home)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Transactions"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        category_column = headers["Category"]
        date_column = headers["Date"]
        year, month_number = (int(part) for part in month.split("-"))
        for row in range(2, sheet.max_row + 1):
            transaction_date = sheet.cell(row=row, column=date_column).value
            if (
                getattr(transaction_date, "year", None) == year
                and getattr(transaction_date, "month", None) == month_number
                and sheet.cell(row=row, column=category_column).value == old
            ):
                sheet.cell(row=row, column=category_column).value = new
                workbook.save(path)
                return
    finally:
        workbook.close()
    raise AssertionError(f"Category {old!r} not found in fixture workbook")


def load_scorecard(home: Path, who: str, month: str = "2026-09"):
    household = load_household(home)
    declarations = load_declarations(home, people=household.slugs)
    plan = load_plan(home, declarations=declarations, people=household.slugs)
    ledger = read_tiller(find_workbook(home / "data", household.tiller_glob))
    return build_scorecard(
        ledger=ledger,
        household=household,
        declarations=declarations,
        plan=plan,
        month=Month.parse(month),
        who=who,
        today=date(2026, 10, 1),
    )


def by_bucket(scorecard) -> dict[str, Any]:
    return {bucket.bucket: bucket for bucket in scorecard.buckets}


def test_scorecard_cli_runs_on_example_instance(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    result = run_hfk(
        "scorecard",
        "--home",
        str(home),
        "--month",
        "2026-09",
        "--who",
        "alex",
        "--today",
        "2026-10-01",
    )
    sam = run_hfk(
        "scorecard",
        "--home",
        str(home),
        "--month",
        "2026-09",
        "--who",
        "sam",
        "--today",
        "2026-10-01",
    )

    assert result.returncode == 0
    assert result.stderr == ""
    assert "scorecard 2026-09 for alex (closed month, shares per declarations)" in result.stdout
    assert "Fixed Essential" in result.stdout
    assert "plan 1200.00  actual 1500.00" in result.stdout
    assert "total  plan 1870.00  actual 1840.00  left 30.00" in result.stdout
    assert sam.returncode == 0
    assert sam.stderr == ""


def test_scorecard_buckets_match_fixture_arithmetic(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    alex = load_scorecard(home, "alex")
    sam = load_scorecard(home, "sam")
    alex_buckets = by_bucket(alex)
    sam_buckets = by_bucket(sam)

    assert tuple(bucket.bucket for bucket in alex.buckets) == (
        "Fixed Essential",
        "Variable Essential",
        "Fixed Discretionary",
        "Variable Discretionary",
    )
    assert alex.plan_lines == 5
    assert alex.transactions == 7
    assert alex.planned_total == Decimal("1870.00")
    assert alex.actual_total == Decimal("1840.00")

    assert alex_buckets["Fixed Essential"].planned == Decimal("1200.00")
    assert alex_buckets["Fixed Essential"].actual == Decimal("1500.00")
    assert alex_buckets["Fixed Essential"].variance == Decimal("-300.00")
    assert alex_buckets["Variable Essential"].planned == Decimal("640.00")
    assert alex_buckets["Variable Essential"].actual == Decimal("287.50")
    assert alex_buckets["Variable Essential"].variance == Decimal("352.50")
    assert alex_buckets["Fixed Discretionary"].planned == Decimal("30.00")
    assert alex_buckets["Fixed Discretionary"].actual == Decimal("0.00")
    assert alex_buckets["Fixed Discretionary"].variance == Decimal("30.00")
    assert alex_buckets["Variable Discretionary"].planned == Decimal("0.00")
    assert alex_buckets["Variable Discretionary"].actual == Decimal("52.50")
    assert alex_buckets["Variable Discretionary"].variance == Decimal("-52.50")

    assert sam.plan_lines == 3
    assert sam.transactions == 6
    assert sam.planned_total == Decimal("230.00")
    assert sam.actual_total == Decimal("340.00")
    assert sam_buckets["Fixed Essential"].planned == Decimal("0.00")
    assert sam_buckets["Fixed Essential"].actual == Decimal("0.00")
    assert sam_buckets["Variable Essential"].planned == Decimal("170.00")
    assert sam_buckets["Variable Essential"].actual == Decimal("287.50")
    assert sam_buckets["Variable Essential"].variance == Decimal("-117.50")
    assert sam_buckets["Fixed Discretionary"].planned == Decimal("0.00")
    assert sam_buckets["Fixed Discretionary"].actual == Decimal("0.00")
    assert sam_buckets["Variable Discretionary"].planned == Decimal("60.00")
    assert sam_buckets["Variable Discretionary"].actual == Decimal("52.50")
    assert sam_buckets["Variable Discretionary"].variance == Decimal("7.50")

    july_alex = by_bucket(load_scorecard(home, "alex", month="2026-07"))
    assert july_alex["Fixed Discretionary"].actual == Decimal("60.00")


def test_both_people_reconcile_to_ledger_spend(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household = load_household(home)
    declarations = load_declarations(home, people=household.slugs)
    ledger = read_tiller(find_workbook(home / "data", household.tiller_glob))
    resolutions = resolve_ledger(ledger, declarations)
    expected = {bucket.name: Decimal("0.00") for bucket in declarations.buckets}

    for transaction in ledger.transactions_in_month(2026, 9):
        resolution = resolutions[transaction.category]
        for allocation in allocate(transaction.amount, resolution):
            expected[allocation.bucket] += allocation.amount

    alex = by_bucket(load_scorecard(home, "alex"))
    sam = by_bucket(load_scorecard(home, "sam"))
    for bucket, amount in expected.items():
        assert alex[bucket].actual + sam[bucket].actual == -amount


def test_unfinished_month_is_refused(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household = load_household(home)
    declarations = load_declarations(home, people=household.slugs)
    plan = load_plan(home, declarations=declarations, people=household.slugs)
    ledger = read_tiller(find_workbook(home / "data", household.tiller_glob))

    with pytest.raises(ScorecardError, match="has not ended"):
        build_scorecard(
            ledger=ledger,
            household=household,
            declarations=declarations,
            plan=plan,
            month=Month.parse("2026-09"),
            who="alex",
            today=date(2026, 9, 30),
        )

    result = run_hfk(
        "scorecard",
        "--home",
        str(home),
        "--month",
        "2026-09",
        "--who",
        "alex",
        "--today",
        "2026-09-30",
    )
    assert result.returncode == 1
    assert "has not ended" in result.stderr


def test_history_start_clamps_the_window(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    household["history_start"] = "2026-09-10"
    dump_yaml(household_path, household)

    scorecard = load_scorecard(home, "alex")
    buckets = by_bucket(scorecard)

    assert scorecard.clamped is True
    assert scorecard.window_start == date(2026, 9, 10)
    assert buckets["Variable Essential"].actual == Decimal("227.50")
    assert buckets["Variable Discretionary"].actual == Decimal("30.00")
    assert buckets["Fixed Essential"].actual == Decimal("0.00")

    household["history_start"] = "2026-10-01"
    dump_yaml(household_path, household)
    with pytest.raises(ScorecardError, match="before history_start"):
        load_scorecard(home, "alex")


def test_plan_lines_effective_in_month(tmp_path: Path) -> None:
    old_home = make_instance(tmp_path / "old", end_month=(2025, 6))
    current_home = make_instance(tmp_path / "current")

    old_scorecard = by_bucket(load_scorecard(old_home, "alex", month="2025-06"))
    current_scorecard = by_bucket(load_scorecard(current_home, "alex", month="2026-09"))

    assert old_scorecard["Variable Essential"].planned == Decimal("590.00")
    assert current_scorecard["Variable Essential"].planned == Decimal("640.00")


def test_who_is_required_with_several_people(tmp_path: Path) -> None:
    home = make_instance(tmp_path / "two")

    result = run_hfk(
        "scorecard",
        "--home",
        str(home),
        "--month",
        "2026-09",
        "--today",
        "2026-10-01",
    )

    assert result.returncode == 1
    assert "choose --who from: alex, sam" in result.stderr

    solo = make_instance(tmp_path / "solo")
    household_path = solo / "household.yaml"
    household = load_yaml(household_path)
    household["people"] = [{"slug": "alex", "name": "Alex"}]
    household["accounts"] = {
        "Alex Checking": {"owner": "alex"},
        "Sam Checking": {"owner": "alex"},
        "Joint Card": {"owner": "alex"},
        "Savings": {"owner": "alex", "kind": "savings"},
    }
    dump_yaml(household_path, household)

    declarations_path = solo / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["sharing"] = []
    dump_yaml(declarations_path, declarations)

    plan_path = solo / "plan.yaml"
    plan = load_yaml(plan_path)
    plan["plan"]["lines"] = [
        line for line in plan["plan"]["lines"] if line.get("who") == "alex"
    ]
    dump_yaml(plan_path, plan)

    solo_result = run_hfk(
        "scorecard",
        "--home",
        str(solo),
        "--month",
        "2026-09",
        "--today",
        "2026-10-01",
    )
    assert solo_result.returncode == 0
    assert "scorecard 2026-09 for alex" in solo_result.stdout


def test_conflicts_stop_the_scorecard(tmp_path: Path) -> None:
    conflict_home = make_instance(tmp_path / "conflict")
    declarations_path = conflict_home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["sharing"].append(
        {
            "match": {"account": "Alex Checking", "description": "City Utilities*"},
            "weights": {"alex": 1.0},
            "reason": "Utilities are kept by Alex in this conflicting example.",
        }
    )
    dump_yaml(declarations_path, declarations)

    result = run_hfk(
        "scorecard",
        "--home",
        str(conflict_home),
        "--month",
        "2026-09",
        "--who",
        "alex",
        "--today",
        "2026-10-01",
    )

    assert result.returncode == 1
    assert "Conflicting sharing rules" in result.stderr

    unmapped_home = make_instance(tmp_path / "unmapped")
    declarations_path = unmapped_home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    del declarations["groups"]["Subscriptions"]
    dump_yaml(declarations_path, declarations)

    result = run_hfk(
        "scorecard",
        "--home",
        str(unmapped_home),
        "--month",
        "2026-09",
        "--who",
        "alex",
        "--today",
        "2026-10-01",
    )

    assert result.returncode == 1
    assert "Subscriptions" in result.stderr
    assert isinstance(result.stderr, str)


def test_cents_tolerance_reconciles(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    append_transaction(
        home,
        day=24,
        description="Tiny Grocer One",
        category="Groceries",
        amount=Decimal("-0.01"),
        account="Joint Card",
    )
    append_transaction(
        home,
        day=25,
        description="Tiny Grocer Three",
        category="Groceries",
        amount=Decimal("-0.03"),
        account="Joint Card",
    )

    household = load_household(home)
    declarations = load_declarations(home, people=household.slugs)
    ledger = read_tiller(find_workbook(home / "data", household.tiller_glob))
    resolutions = resolve_ledger(ledger, declarations)
    ledger_by_bucket = {bucket.name: Decimal("0.00") for bucket in declarations.buckets}
    person_by_bucket = {
        slug: {bucket.name: Decimal("0.00") for bucket in declarations.buckets}
        for slug in household.slugs
    }

    for transaction in ledger.transactions_in_month(2026, 9):
        resolution = resolutions[transaction.category]
        for allocation in allocate(transaction.amount, resolution):
            ledger_by_bucket[allocation.bucket] += allocation.amount
        shares = share_transaction(transaction, household, declarations)
        for slug in household.slugs:
            for allocation in allocate(shares.amount_for(slug), resolution):
                person_by_bucket[slug][allocation.bucket] += allocation.amount

    alex = by_bucket(load_scorecard(home, "alex"))
    sam = by_bucket(load_scorecard(home, "sam"))
    for bucket, amount in ledger_by_bucket.items():
        assert alex[bucket].actual + sam[bucket].actual == -amount
        assert abs(alex[bucket].actual + person_by_bucket["alex"][bucket]) <= Decimal("0.01")
        assert abs(sam[bucket].actual + person_by_bucket["sam"][bucket]) <= Decimal("0.01")


def test_non_spend_is_excluded(tmp_path: Path) -> None:
    home = make_instance(tmp_path / "income-transfer")
    base_alex = by_bucket(load_scorecard(home, "alex"))
    append_transaction(
        home,
        day=26,
        description="Extra Example Paycheck",
        category="Alex Paycheck",
        amount=Decimal("9999.00"),
        account="Alex Checking",
    )
    append_transaction(
        home,
        day=27,
        description="Extra Savings Transfer",
        category="Savings Transfer",
        amount=Decimal("-9999.00"),
        account="Alex Checking",
    )
    with_non_spend = by_bucket(load_scorecard(home, "alex"))
    assert with_non_spend == base_alex

    excluded_home = make_instance(tmp_path / "excluded-group")
    declarations_path = excluded_home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["non_spend_groups"].append("Food")
    dump_yaml(declarations_path, declarations)
    excluded = by_bucket(load_scorecard(excluded_home, "alex"))
    assert excluded["Variable Essential"].actual == Decimal("105.00")

    refund_home = make_instance(tmp_path / "refund")
    append_transaction(
        refund_home,
        day=28,
        description="Grocer Refund",
        category="Groceries",
        amount=Decimal("30.00"),
        account="Joint Card",
    )
    refunded = by_bucket(load_scorecard(refund_home, "alex"))
    assert refunded["Variable Essential"].actual == Decimal("272.50")


def test_history_start_window_rules(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    del household["history_start"]
    dump_yaml(household_path, household)

    unclamped = load_scorecard(home, "alex")
    assert unclamped.clamped is False
    assert unclamped.window_start == date(2026, 9, 1)

    household["history_start"] = "2026-09-10"
    dump_yaml(household_path, household)
    clamped = load_scorecard(home, "alex")
    buckets = by_bucket(clamped)
    assert clamped.clamped is True
    assert clamped.window_start == date(2026, 9, 10)
    assert buckets["Variable Essential"].actual == Decimal("227.50")
    assert buckets["Fixed Essential"].actual == Decimal("0.00")

    household["history_start"] = "2026-10-01"
    dump_yaml(household_path, household)
    with pytest.raises(ScorecardError, match="before history_start"):
        load_scorecard(home, "alex")


def test_renamed_category_is_never_remapped(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    rename_first_transaction_category(home, "Groceries", "Groceries (old)")

    with pytest.raises(BucketError, match=r"Groceries \(old\)"):
        load_scorecard(home, "alex")

    result = run_hfk(
        "scorecard",
        "--home",
        str(home),
        "--month",
        "2026-09",
        "--who",
        "alex",
        "--today",
        "2026-10-01",
    )
    assert result.returncode == 1
    assert "Groceries (old)" in result.stderr

    declarations_path = home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["overrides"]["Groceries (old)"] = {
        "bucket": "Variable Discretionary",
        "reason": "Renamed example category kept separate for history.",
    }
    dump_yaml(declarations_path, declarations)

    buckets = by_bucket(load_scorecard(home, "alex"))
    assert buckets["Variable Essential"].actual == Decimal("227.50")
    assert buckets["Variable Discretionary"].actual == Decimal("112.50")


def test_lines_without_categories_count_and_paced_is_ignored(tmp_path: Path) -> None:
    unpaced_home = make_instance(tmp_path / "unpaced")
    paced_home = make_instance(tmp_path / "paced")
    for home, paced in ((unpaced_home, False), (paced_home, True)):
        plan_path = home / "plan.yaml"
        plan = load_yaml(plan_path)
        streaming = plan["plan"]["lines"][6]
        assert "categories" not in streaming
        streaming["monthly"] = 25.00
        if paced:
            streaming["paced"] = True
        dump_yaml(plan_path, plan)

    unpaced = load_scorecard(unpaced_home, "alex")
    paced = load_scorecard(paced_home, "alex")
    unpaced_buckets = by_bucket(unpaced)
    paced_buckets = by_bucket(paced)

    assert unpaced_buckets["Fixed Discretionary"].planned == Decimal("35.00")
    assert paced_buckets["Fixed Discretionary"].planned == Decimal("35.00")
    assert unpaced.buckets == paced.buckets
