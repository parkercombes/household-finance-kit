from __future__ import annotations

import os
import re
import subprocess
from pathlib import Path

from openpyxl import load_workbook

from hfk.source.fixture import build_fixture

ROOT = Path(__file__).resolve().parents[1]
PYTHONPATH = str(ROOT / "src")
PYTHON = (
    "/private/tmp/claude-501/-Users-bpc/f9f8878c-626f-44b9-a38e-62958c402f3d/"
    "scratchpad/venv/bin/python"
)
SCHEMA = ROOT / "docs" / "TILLER-SCHEMA.md"


def read_rows(path: Path, sheet_name: str) -> tuple[tuple[object, ...], ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    return tuple(tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True))


def headers(path: Path, sheet_name: str) -> list[str]:
    return [value for value in read_rows(path, sheet_name)[0] if value is not None]


def schema_headers(section: str) -> list[str]:
    lines = SCHEMA.read_text(encoding="utf-8").splitlines()
    in_section = False
    found_table = False
    result: list[str] = []
    for line in lines:
        if line == f"## `{section}`":
            in_section = True
            continue
        if in_section and line.startswith("## "):
            break
        if not in_section or not line.startswith("|"):
            continue
        found_table = True
        cells = [cell.strip().strip("`") for cell in line.strip().strip("|").split("|")]
        if cells[0] in {"Header", "---", "Monthly columns"}:
            continue
        result.append(cells[0])
    assert found_table
    return result


def row_maps(path: Path, sheet_name: str) -> list[dict[str, object]]:
    rows = read_rows(path, sheet_name)
    header = list(rows[0])
    return [
        {str(name): value for name, value in zip(header, row, strict=True)}
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def cell_values(path: Path) -> tuple[tuple[str, tuple[tuple[object, ...], ...]], ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return tuple(
        (sheet_name, tuple(tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)))
        for sheet_name in workbook.sheetnames
    )


def test_sheet_names(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)

    assert workbook.sheetnames == ["Transactions", "Categories", "Balance History", "Accounts"]


def test_headers_match_schema_tables(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")

    assert headers(path, "Transactions") == schema_headers("Transactions")
    assert headers(path, "Categories") == schema_headers("Categories")
    assert headers(path, "Balance History") == schema_headers("Balance History")


def test_determinism(tmp_path: Path) -> None:
    first = build_fixture(tmp_path / "first.xlsx")
    second = build_fixture(tmp_path / "second.xlsx")

    assert cell_values(first) == cell_values(second)


def test_shuffle_changes_order_not_header_set(tmp_path: Path) -> None:
    plain = build_fixture(tmp_path / "plain.xlsx")
    shuffled = build_fixture(tmp_path / "shuffled.xlsx", shuffle_headers=True)

    assert headers(shuffled, "Transactions") != headers(plain, "Transactions")
    assert set(headers(shuffled, "Transactions")) == set(headers(plain, "Transactions"))


def test_drop_columns_removes_header(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx", drop_columns=[("Transactions", "Account #")])

    assert "Account #" not in headers(path, "Transactions")


def test_extra_columns_are_added(tmp_path: Path) -> None:
    path = build_fixture(
        tmp_path / "fixture.xlsx",
        extra_columns=[("Transactions", "Fixture Extra"), ("Balance History", "Review Note")],
    )

    assert "Fixture Extra" in headers(path, "Transactions")
    assert "Review Note" in headers(path, "Balance History")


def test_expenses_negative_income_positive(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    categories = {row["Category"]: row["Type"] for row in row_maps(path, "Categories")}
    transactions = row_maps(path, "Transactions")

    expense_amounts = [
        row["Amount"] for row in transactions if categories[row["Category"]] == "Expense"
    ]
    income_amounts = [row["Amount"] for row in transactions if categories[row["Category"]] == "Income"]

    assert expense_amounts
    assert income_amounts
    assert all(amount < 0 for amount in expense_amounts)
    assert all(amount > 0 for amount in income_amounts)


def test_exactly_one_empty_balance_when_enabled(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx", balance_gap=True)
    balances = row_maps(path, "Balance History")

    assert sum(row["Balance"] is None for row in balances) == 1


def test_no_empty_balance_when_disabled(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx", balance_gap=False)
    balances = row_maps(path, "Balance History")

    assert all(row["Balance"] is not None for row in balances)


def test_no_real_looking_email_or_long_digit_run(tmp_path: Path) -> None:
    path = build_fixture(tmp_path / "fixture.xlsx")
    email = re.compile(r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b")
    long_digits = re.compile(r"\d{9,}")

    for _sheet_name, rows in cell_values(path):
        for row in rows:
            for value in row:
                if isinstance(value, str):
                    assert email.search(value) is None
                    assert long_digits.search(value) is None


def test_cli_dev_fixture_end_to_end(tmp_path: Path) -> None:
    out = tmp_path / "f.xlsx"
    env = os.environ.copy()
    env["PYTHONPATH"] = PYTHONPATH
    result = subprocess.run(
        [PYTHON, "-m", "hfk.cli", "dev", "fixture", "--out", str(out)],
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0
    assert result.stdout.strip() == str(out)
    assert result.stderr == ""
    assert out.exists()
    assert load_workbook(out, read_only=True, data_only=True).sheetnames == [
        "Transactions",
        "Categories",
        "Balance History",
        "Accounts",
    ]
