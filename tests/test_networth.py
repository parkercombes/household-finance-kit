"""Net worth engine and CLI tests (M4-03)."""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from hfk.config import load_declarations, load_household
from hfk.engine import NetWorthError, build_networth
from hfk.source import find_workbook, read_tiller
from hfk.source.fixture import build_fixture

ROOT = Path(__file__).resolve().parents[1]
INSTANCE_EXAMPLE = ROOT / "examples" / "instance"
PYTHONPATH = str(ROOT / "src")


def make_instance(tmp_path: Path, **fixture_kwargs: Any) -> Path:
    tmp_path.mkdir(parents=True, exist_ok=True)
    for filename in ("household.yaml", "declarations.yaml", "plan.yaml"):
        shutil.copyfile(INSTANCE_EXAMPLE / filename, tmp_path / filename)
    build_fixture(tmp_path / "data" / "Tiller-Example.xlsx", **fixture_kwargs)
    return tmp_path


def run_hfk(*args: str) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONPATH"] = PYTHONPATH
    return subprocess.run(
        [sys.executable, "-m", "hfk", *args],
        cwd=ROOT,
        env=env,
        text=True,
        capture_output=True,
        check=False,
    )


def load_yaml(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    return data


def dump_yaml(path: Path, data: dict[str, Any]) -> None:
    path.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")


def fixture_workbook(home: Path) -> Path:
    return next((home / "data").glob("Tiller*.xlsx"))


def append_balance(
    home: Path,
    *,
    observed: date,
    account: str,
    balance: Decimal | None,
    account_class: str = "Asset",
    account_type: str = "Checking",
) -> None:
    path = fixture_workbook(home)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Balance History"]
        headers = [cell.value for cell in sheet[1]]
        account_id = account.lower().replace(" ", "-")
        values = {
            "Date": observed,
            "Time": None,
            "Account": account,
            "Account #": None,
            "Account ID": f"acct-{account_id}",
            "Balance ID": f"test-{observed.isoformat()}-{account_id}",
            "Institution": "Example Bank",
            "Balance": balance,
            "Month": observed.replace(day=1),
            "Week": observed,
            "Type": account_type,
            "Class": account_class,
            "Account Status": "Open",
            "Date Added": observed,
        }
        sheet.append([values.get(header) for header in headers])
        workbook.save(path)
    finally:
        workbook.close()


def blank_account_class(home: Path, account: str) -> None:
    path = fixture_workbook(home)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Balance History"]
        headers = [cell.value for cell in sheet[1]]
        account_at = headers.index("Account")
        class_at = headers.index("Class")
        for row in sheet.iter_rows(min_row=2):
            if row[account_at].value == account:
                row[class_at].value = None
        workbook.save(path)
    finally:
        workbook.close()


def load_networth(home: Path, as_of: date, stale_after_days: int = 14):
    household = load_household(home)
    declarations = load_declarations(home, people=household.slugs)
    ledger = read_tiller(find_workbook(home / "data", household.tiller_glob))
    return build_networth(
        ledger=ledger,
        household=household,
        declarations=declarations,
        as_of=as_of,
        stale_after_days=stale_after_days,
    )


def by_account(report) -> dict[str, Any]:
    return {line.account: line for line in report.lines}


def test_networth_cli_runs_on_example_instance(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")

    assert result.returncode == 0
    assert result.stderr == ""
    lines = result.stdout.splitlines()
    assert lines[0] == "net worth as of 2026-09-30"
    assert f"{'Alex Checking':<13}  {'Asset':<9}  3025.00  observed 2026-09-28" in lines
    assert f"{'Sam Checking':<13}  {'Asset':<9}  2725.00  observed 2026-09-28" in lines
    assert (
        f"{'Joint Card':<13}  {'Liability':<9}  25.00  "
        "observed 2026-09-28; liability positive"
    ) in lines
    assert f"{'Savings':<13}  {'Asset':<9}  5625.00  observed 2026-09-28" in lines
    assert "assets 11375.00" in lines
    assert "liabilities 25.00" in lines
    assert "net worth 11400.00" in lines
    assert "counted 4 of 4  excluded 0" in lines
    assert not any(line.startswith("warning:") for line in lines)
    assert not any(line.startswith("note:") for line in lines)

    same_day = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-28")
    assert same_day.returncode == 0
    assert same_day.stderr == ""
    same_day_lines = same_day.stdout.splitlines()
    assert same_day_lines[0] == "net worth as of 2026-09-28"
    assert "net worth 11400.00" in same_day_lines
    assert "counted 4 of 4  excluded 0" in same_day_lines

    july = run_hfk("networth", "--home", str(home), "--as-of", "2026-07-31")
    assert july.returncode == 0
    assert july.stderr == ""
    july_lines = july.stdout.splitlines()
    assert "assets 10625.00" in july_lines
    assert "liabilities -225.00" in july_lines
    assert "net worth 10400.00" in july_lines


def test_gap_is_excluded_never_zero(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    report = load_networth(home, date(2026, 6, 15))
    joint = by_account(report)["Joint Card"]

    assert joint.gap is True
    assert joint.balance is None
    assert joint.observed == date(2026, 5, 28)
    assert joint.gap_reason == "latest observation 2026-05-28 is blank"
    assert joint.counted is False
    assert joint.flags == ()
    assert report.assets == Decimal("9875.00")
    assert report.liabilities == Decimal("0.00")
    assert report.net_worth == Decimal("9875.00")
    assert report.counted == 3
    assert report.excluded == ("Joint Card",)
    assert report.incomplete == ("Joint Card",)
    for name in ("Alex Checking", "Sam Checking", "Savings"):
        assert by_account(report)[name].flags == ("stale (18 days)",)

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-06-15")
    assert result.returncode == 0
    assert result.stderr == ""
    lines = result.stdout.splitlines()
    assert lines[0] == "net worth as of 2026-06-15"
    assert (
        f"{'Joint Card':<13}  {'Liability':<9}  gap  "
        "latest observation 2026-05-28 is blank; excluded"
    ) in lines
    assert (
        f"{'Alex Checking':<13}  {'Asset':<9}  2525.00  "
        "observed 2026-05-28; stale (18 days)"
    ) in lines
    assert "assets 9875.00" in lines
    assert "liabilities 0.00" in lines
    assert "net worth 9875.00" in lines
    assert "counted 3 of 4  excluded 1" in lines
    assert "warning: incomplete net worth; excluded: Joint Card" in lines


def test_classes_split_assets_and_liabilities(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    report = load_networth(home, date(2026, 9, 30))
    assert [line.account for line in report.lines] == [
        "Alex Checking",
        "Sam Checking",
        "Joint Card",
        "Savings",
    ]
    assert by_account(report)["Alex Checking"].account_class == "Asset"
    assert by_account(report)["Joint Card"].account_class == "Liability"
    assert report.assets == Decimal("11375.00")
    assert report.liabilities == Decimal("25.00")
    assert report.net_worth == Decimal("11400.00")

    july = load_networth(home, date(2026, 7, 31))
    joint = by_account(july)["Joint Card"]
    assert joint.balance == Decimal("-225.00")
    assert joint.observed == date(2026, 7, 28)
    assert joint.counted is True
    assert joint.flags == ()
    assert july.assets == Decimal("10625.00")
    assert july.liabilities == Decimal("-225.00")
    assert july.net_worth == Decimal("10400.00")

    append_balance(
        home,
        observed=date(2026, 9, 20),
        account="House Fund",
        balance=Decimal("4100.00"),
        account_class=" Asset ",
        account_type="Savings",
    )
    normalized = load_networth(home, date(2026, 9, 30))
    house = by_account(normalized)["House Fund"]
    assert house.account_class == " Asset "
    assert house.counted is True
    assert house.flags == ("undeclared",)
    assert normalized.assets == Decimal("15475.00")
    assert normalized.net_worth == Decimal("15500.00")
    assert normalized.counted == 5
    assert normalized.excluded == ()


def test_positive_liability_is_flagged(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    report = load_networth(home, date(2026, 9, 30))
    joint = by_account(report)["Joint Card"]
    assert joint.account_class == "Liability"
    assert joint.balance == Decimal("25.00")
    assert joint.counted is True
    assert joint.flags == ("liability positive",)
    assert report.liabilities == Decimal("25.00")

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")
    assert result.returncode == 0
    assert (
        f"{'Joint Card':<13}  {'Liability':<9}  25.00  "
        "observed 2026-09-28; liability positive"
    ) in result.stdout.splitlines()


def test_stale_synced_account_is_flagged(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    report = load_networth(home, date(2026, 10, 20))
    assert by_account(report)["Alex Checking"].flags == ("stale (22 days)",)
    assert by_account(report)["Sam Checking"].flags == ("stale (22 days)",)
    assert by_account(report)["Joint Card"].flags == ("liability positive", "stale (22 days)")
    assert by_account(report)["Savings"].flags == ("stale (22 days)",)
    assert report.counted == 4
    assert report.excluded == ()
    assert report.net_worth == Decimal("11400.00")

    relaxed = load_networth(home, date(2026, 10, 20), stale_after_days=30)
    assert by_account(relaxed)["Alex Checking"].flags == ()
    assert by_account(relaxed)["Joint Card"].flags == ("liability positive",)
    assert relaxed.stale_after_days == 30
    assert relaxed.net_worth == Decimal("11400.00")

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-10-20")
    assert result.returncode == 0
    assert (
        f"{'Alex Checking':<13}  {'Asset':<9}  3025.00  "
        "observed 2026-09-28; stale (22 days)"
    ) in result.stdout.splitlines()

    relaxed_result = run_hfk(
        "networth", "--home", str(home), "--as-of", "2026-10-20", "--stale-after-days", "30"
    )
    assert relaxed_result.returncode == 0
    assert "stale (" not in relaxed_result.stdout

    with pytest.raises(NetWorthError) as error:
        load_networth(home, date(2026, 10, 20), stale_after_days=-1)
    assert "stale_after_days must be a non-negative integer" in str(error.value)

    refused = run_hfk(
        "networth", "--home", str(home), "--as-of", "2026-10-20", "--stale-after-days", "-1"
    )
    assert refused.returncode == 1
    assert refused.stderr.startswith("networth: ")
    assert "stale_after_days" in refused.stderr


def test_unsynced_account_is_never_stale(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    household["accounts"]["Savings"]["unsynced"] = True
    dump_yaml(household_path, household)

    report = load_networth(home, date(2026, 10, 20))
    savings = by_account(report)["Savings"]
    assert savings.unsynced is True
    assert savings.counted is True
    assert savings.flags == ("unsynced",)
    assert savings.balance == Decimal("5625.00")
    assert by_account(report)["Alex Checking"].flags == ("stale (22 days)",)
    assert by_account(report)["Joint Card"].flags == ("liability positive", "stale (22 days)")
    assert report.assets == Decimal("11375.00")
    assert report.net_worth == Decimal("11400.00")

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-10-20")
    assert result.returncode == 0
    assert result.stderr == ""
    assert (
        f"{'Savings':<13}  {'Asset':<9}  5625.00  observed 2026-09-28; unsynced"
    ) in result.stdout.splitlines()

    household["accounts"]["Savings"]["unsynced"] = "yes"
    dump_yaml(household_path, household)
    rejected = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")
    assert rejected.returncode == 1
    assert rejected.stdout == ""
    assert rejected.stderr.startswith("networth: ")
    assert "unsynced" in rejected.stderr
    assert "accounts.Savings.unsynced" in rejected.stderr


def test_closed_account_is_excluded_after_close(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    del household["accounts"]["Savings"]
    dump_yaml(household_path, household)
    declarations_path = home / "declarations.yaml"
    declarations = load_yaml(declarations_path)
    declarations["closed_accounts"] = {"Savings": {"closed": "2026-07-01"}}
    dump_yaml(declarations_path, declarations)

    report = load_networth(home, date(2026, 9, 30))
    assert [line.account for line in report.lines] == [
        "Alex Checking",
        "Sam Checking",
        "Joint Card",
        "Savings",
    ]
    savings = by_account(report)["Savings"]
    assert savings.balance == Decimal("5250.00")
    assert savings.observed == date(2026, 6, 28)
    assert savings.closed == date(2026, 7, 1)
    assert savings.counted is False
    assert savings.flags == ("closed 2026-07-01",)
    assert savings.unsynced is False
    assert report.excluded == ("Savings",)
    assert report.incomplete == ()
    assert report.assets == Decimal("5750.00")
    assert report.liabilities == Decimal("25.00")
    assert report.net_worth == Decimal("5775.00")
    assert report.counted == 3

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")
    assert result.returncode == 0
    assert result.stderr == ""
    lines = result.stdout.splitlines()
    assert (
        f"{'Savings':<13}  {'Asset':<9}  5250.00  "
        "observed 2026-06-28; closed 2026-07-01; excluded"
    ) in lines
    assert "net worth 5775.00" in lines
    assert "counted 3 of 4  excluded 1" in lines
    assert not any(line.startswith("warning:") for line in lines)

    still_open = load_networth(home, date(2026, 6, 15))
    june = by_account(still_open)["Savings"]
    assert june.counted is True
    assert june.balance == Decimal("5125.00")
    assert june.flags == ("stale (18 days)",)
    assert still_open.assets == Decimal("9875.00")
    assert still_open.liabilities == Decimal("0.00")
    assert still_open.net_worth == Decimal("9875.00")
    assert still_open.excluded == ("Joint Card",)
    assert still_open.incomplete == ("Joint Card",)


def test_unclassified_account_is_excluded(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    blank_account_class(home, "Sam Checking")

    report = load_networth(home, date(2026, 9, 30))
    sam = by_account(report)["Sam Checking"]
    assert sam.account_class is None
    assert sam.balance == Decimal("2725.00")
    assert sam.observed == date(2026, 9, 28)
    assert sam.counted is False
    assert sam.flags == ("unclassified",)
    assert report.excluded == ("Sam Checking",)
    assert report.incomplete == ("Sam Checking",)
    assert report.assets == Decimal("8650.00")
    assert report.liabilities == Decimal("25.00")
    assert report.net_worth == Decimal("8675.00")
    assert report.counted == 3

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")
    assert result.returncode == 0
    assert result.stderr == ""
    lines = result.stdout.splitlines()
    assert (
        f"{'Sam Checking':<13}  {'-':<9}  2725.00  "
        "observed 2026-09-28; unclassified; excluded"
    ) in lines
    assert "net worth 8675.00" in lines
    assert "counted 3 of 4  excluded 1" in lines
    assert "warning: incomplete net worth; excluded: Sam Checking" in lines


def test_no_observations_is_refused(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    with pytest.raises(NetWorthError) as error:
        load_networth(home, date(2026, 3, 31))
    assert str(error.value) == (
        "no account has a balance observation on or before 2026-03-31. "
        "Fix: choose a later date or check the Balance History sheet"
    )

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-03-31")
    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("networth: ")
    assert "no account has a balance observation on or before 2026-03-31" in result.stderr


def test_as_of_before_history_start_is_refused(tmp_path: Path) -> None:
    home = make_instance(tmp_path)

    with pytest.raises(NetWorthError) as error:
        load_networth(home, date(2024, 12, 31))
    message = str(error.value)
    assert "as_of 2024-12-31 is before history_start 2025-01-01" in message
    assert "choose a date on or after history_start" in message

    result = run_hfk("networth", "--home", str(home), "--as-of", "2024-12-31")
    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("networth: ")
    assert "before history_start 2025-01-01" in result.stderr


def test_undeclared_account_is_flagged_and_still_counted(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    del household["accounts"]["Savings"]
    dump_yaml(household_path, household)

    report = load_networth(home, date(2026, 9, 30))
    savings = by_account(report)["Savings"]
    assert savings.declared is False
    assert savings.counted is True
    assert savings.flags == ("undeclared",)
    assert report.excluded == ()
    assert report.incomplete == ()
    assert report.net_worth == Decimal("11400.00")
    assert report.counted == 4

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")
    assert result.returncode == 0
    assert (
        f"{'Savings':<13}  {'Asset':<9}  5625.00  observed 2026-09-28; undeclared"
    ) in result.stdout.splitlines()


def test_observations_before_history_start_are_noted(tmp_path: Path) -> None:
    home = make_instance(tmp_path)
    household_path = home / "household.yaml"
    household = load_yaml(household_path)
    household["history_start"] = "2026-06-01"
    dump_yaml(household_path, household)

    report = load_networth(home, date(2026, 9, 30))
    assert report.ignored_before_history_start == 8
    assert report.counted == 4
    assert by_account(report)["Joint Card"].balance == Decimal("25.00")
    assert report.net_worth == Decimal("11400.00")

    result = run_hfk("networth", "--home", str(home), "--as-of", "2026-09-30")
    assert result.returncode == 0
    assert (
        "note: ignored 8 observations before history_start 2026-06-01"
        in result.stdout.splitlines()
    )
